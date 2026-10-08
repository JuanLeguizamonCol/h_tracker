"""Excel export of a project performance review — reproduces the Impact Point
review workbook layout (columns Criteria / Sub-Criteria / Score / Reviewer
Notes; blue = employee sections, green = reviewer sections)."""
import math
import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from services.performance_reviews import REVIEW_TEMPLATE

NAVY = "FF0B347D"         # employee section headers
GREEN = "FF237F2E"        # reviewer section headers
BLACK = "FF000000"        # "Criteria / Average Score" summary header
LIGHT_BLUE = "FFB9DDFC"   # employee field labels
LIGHT_GREEN = "FFC6E7C8"  # criterion rows
TAN = "FFD0AF8F"          # summary criterion labels
CREAM = "FFFFF3DF"        # summary criterion averages
PEACH = "FFFFE1AF"        # summary overall average
WHITE = "FFFFFFFF"

COL_WIDTHS = {"A": 42.19, "B": 73.13, "C": 32.66, "D": 83.44}

H12 = Font(name="Calibri", bold=True, size=12)
H12_WHITE = Font(name="Calibri", bold=True, size=12, color=WHITE)
BOLD10 = Font(name="Calibri", bold=True, size=10)
BODY10 = Font(name="Calibri", size=10)

LEFT = Alignment(horizontal="left", vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _fill(argb: str) -> PatternFill:
    return PatternFill("solid", fgColor=argb)


def _cell(ws, ref: str, value, font=BODY10, fill=None, align=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if fill:
        c.fill = _fill(fill)
    if align:
        c.alignment = align
    return c


def _fit_row(ws, row: int) -> None:
    """openpyxl can't auto-fit; estimate a height from the wrapped text so long
    notes aren't clipped to a single line when the file is opened."""
    lines = 1
    for col, width in COL_WIDTHS.items():
        value = ws[f"{col}{row}"].value
        if not isinstance(value, str) or not value:
            continue
        chars_per_line = max(int(width * 1.15), 10)
        n = sum(max(1, math.ceil(len(part) / chars_per_line)) for part in value.split("\n"))
        lines = max(lines, n)
    if lines > 1:
        ws.row_dimensions[row].height = lines * 13.5 + 2


def _hours_label(hours) -> str:
    if hours is None:
        return ""
    return f"{hours:g} hours"


def export_filename(data: dict) -> str:
    who = data.get("client_name") or data["project_name"]
    name = f"{who} - {data['employee_name']} - {data['review_date'].year}.xlsx"
    return re.sub(r'[\\/:*?"<>|]', "", name)


def generate_performance_review_xlsx(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    title = f"{data.get('client_name') or data['project_name']} - {data['employee_name']}"
    ws.title = re.sub(r"[\\/*?:\[\]]", "", title)[:31]
    for col, width in COL_WIDTHS.items():
        ws.column_dimensions[col].width = width

    # Header + instructions
    for col, label in zip("ABCD", ("Criteria", "Sub-Criteria", "Score", "Reviewer Notes")):
        _cell(ws, f"{col}1", label, H12)
    _cell(ws, "A2", "Instructions:", BOLD10, align=LEFT)
    _cell(ws, "B2", "Employees complete sections in Blue.")
    _cell(ws, "B3", "Reviewers complete sections in Green.")

    # Project details (left) + average score summary (right)
    _cell(ws, "A5", "Project Details", H12_WHITE, NAVY, LEFT)
    _cell(ws, "B5", None, H12_WHITE, NAVY)
    _cell(ws, "C5", "Criteria", H12_WHITE, BLACK, LEFT)
    _cell(ws, "D5", "Average Score", H12_WHITE, BLACK)
    _cell(ws, "A6", "Review for:", align=LEFT)
    _cell(ws, "B6", data.get("employee_display_name") or data["employee_name"])

    details = [
        ("Project Name", data.get("client_name") or data["project_name"]),
        ("Reviewer", data.get("reviewer_name") or ""),
        ("Project Description", data.get("project_description") or ""),
        ("Employee's Role / Key Activities", data.get("employee_role") or ""),
        ("Duration", _hours_label(data.get("duration_hours"))),
        ("Date of Review", data["review_date"].strftime("%m/%d/%Y")),
    ]
    summary = [(c["label"], c["average"]) for c in data["criteria_averages"]]
    summary.append(("Average Total", data.get("overall_average")))
    for i, (label, value) in enumerate(details):
        row = 7 + i
        _cell(ws, f"A{row}", label, BOLD10, LIGHT_BLUE, LEFT)
        _cell(ws, f"B{row}", value, align=WRAP)
        s_label, s_value = summary[i]
        _cell(ws, f"C{row}", s_label, BOLD10, TAN, CENTER)
        _cell(ws, f"D{row}", s_value, BOLD10, PEACH if s_label == "Average Total" else CREAM, CENTER)
        _fit_row(ws, row)

    # Self assessment (employee) + reviewer notes on each row
    _cell(ws, "A13", "Self Assessment", H12_WHITE, NAVY, LEFT)
    _cell(ws, "B13", "Self-Assessment Notes", H12_WHITE, NAVY)
    _cell(ws, "C13", None, H12_WHITE, NAVY)
    _cell(ws, "D13", "Reviewer Notes", H12_WHITE, GREEN)
    self_rows = [
        ("Strengths", "self_strengths", "reviewer_strengths_notes"),
        ("Areas for Improvement", "self_improvement", "reviewer_improvement_notes"),
        ("Personal Development Needs", "self_development", "reviewer_development_notes"),
    ]
    for i, (label, self_key, reviewer_key) in enumerate(self_rows):
        row = 14 + i
        _cell(ws, f"A{row}", label, BOLD10, LIGHT_BLUE, LEFT)
        _cell(ws, f"B{row}", data.get(self_key) or None, align=WRAP)
        _cell(ws, f"D{row}", data.get(reviewer_key) or None, align=WRAP)
        _fit_row(ws, row)

    # Reviewer assessment
    for col, label in zip("ABCD", ("Reviewer Assessment", "Sub-Criteria", "Score", "Reviewer Notes")):
        _cell(ws, f"{col}18", label, H12_WHITE, GREEN, CENTER if col == "C" else LEFT)

    scores = data.get("scores") or {}
    avg_by_key = {c["key"]: c["average"] for c in data["criteria_averages"]}
    row = 19
    for criterion in REVIEW_TEMPLATE:
        _cell(ws, f"A{row}", criterion["label"], BOLD10, LIGHT_GREEN, LEFT)
        _cell(ws, f"B{row}", None, BODY10, LIGHT_GREEN)
        _cell(ws, f"C{row}", avg_by_key.get(criterion["key"]), BODY10, LIGHT_GREEN, CENTER_WRAP)
        _cell(ws, f"D{row}", None, BODY10, LIGHT_GREEN)
        row += 1
        for sub_key, sub_label in criterion["sub_criteria"]:
            entry = scores.get(sub_key) or {}
            _cell(ws, f"B{row}", sub_label, align=WRAP)
            _cell(ws, f"C{row}", entry.get("score"), align=CENTER_WRAP)
            _cell(ws, f"D{row}", entry.get("notes"), align=WRAP)
            _fit_row(ws, row)
            row += 1

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
