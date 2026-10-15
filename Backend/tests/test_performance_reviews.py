"""Project performance reviews — services/performance_reviews.py and the
.xlsx export. The sample numbers are the THERAPAK / Michael Franz 2025 review
workbook this format was copied from (category averages 3 / 3.25 / 2.8 / 3 /
2.75, overall 2.96)."""
from datetime import date, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from config.database import Base
import models  # noqa: F401 - registers every model on Base.metadata
from models.clients import Client
from models.employees import Employee
from models.projects import Project
from models.time_entries import TimeEntry
from schemas.performance_reviews import PerformanceReviewCreate
from services import performance_reviews as svc
from services.export_performance_review import generate_performance_review_xlsx, export_filename

SAMPLE_SCORES = {
    "written_2": 3, "written_3": 3, "written_4": 3,
    "verbal_1": 4, "verbal_2": 3, "verbal_3": 3, "verbal_5": 3,
    "analysis_1": 2, "analysis_2": 3, "analysis_3": 3, "analysis_4": 3, "analysis_5": 3,
    "client_1": 2, "client_2": 3, "client_3": 4, "client_4": 3,
    "professionalism_1": 2, "professionalism_2": 3, "professionalism_3": 3, "professionalism_4": 3,
}


@pytest.fixture
def db():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    session.add_all([
        Employee(id="emp", user_id="emp", name="Michael Franz", email="m@x.com", first_name="Michael", last_name="Franz"),
        Employee(id="mgr", user_id="mgr", name="Brandon Beal", email="b@x.com"),
        Employee(id="mgr2", user_id="mgr2", name="Casey Diaz", email="c@x.com"),
        Client(id="cli", name="THERAPAK, LLC"),
    ])
    session.flush()
    session.add(Project(id="prj", client_id="cli", name="THERAPAK LLC", manager_id="mgr", description="Cash flow model"))
    session.flush()
    session.add_all([
        TimeEntry(id="t1", user_id="emp", project_id="prj", date=date(2025, 3, 1), hours=200),
        TimeEntry(id="t2", user_id="emp", project_id="prj", date=date(2025, 9, 1), hours=96),
        TimeEntry(id="t3", user_id="mgr", project_id="prj", date=date(2025, 9, 1), hours=50),
    ])
    session.commit()
    yield session
    session.close()


def _scores(raw):
    return {k: {"score": v, "notes": None} for k, v in raw.items()}


def _reviewer_ids(review):
    return [r.reviewer_id for r in review.reviewers]


def _submit_self_and_reviewer(db, review, reviewer_id="mgr", reviewer_scores=None):
    """Helper: self-assessment then one reviewer's evaluation, both submitted."""
    svc.transition_review(db, review, "submit_self")
    svc.apply_reviewer_evaluation(db, review, reviewer_id, {"scores": _scores(reviewer_scores or {"written_1": 3})})
    db.commit()
    return svc.transition_review(db, review, "submit_manager", actor_id=reviewer_id)


def test_averages_match_the_sample_workbook():
    criteria, overall = svc.compute_averages(_scores(SAMPLE_SCORES))
    assert [c["average"] for c in criteria] == [3.0, 3.25, 2.8, 3.0, 2.75]
    assert overall == 2.96


def test_unrated_criteria_are_skipped_not_zeroed():
    criteria, overall = svc.compute_averages(_scores({"written_1": 4, "verbal_1": 2}))
    assert [c["average"] for c in criteria] == [4.0, 2.0, None, None, None]
    assert overall == 3.0
    assert svc.compute_averages({}) == ([{**c, "average": None} for c in criteria], None)


def test_normalize_scores_validates():
    with pytest.raises(ValueError):
        svc.normalize_scores({"nope": {"score": 3}})
    with pytest.raises(ValueError):
        svc.normalize_scores({"written_1": {"score": 6}})
    # A bare {score: None, notes: None} is kept, not dropped — it's an
    # explicit N/A (see ScorePicker/naSelected in PerformanceReviewPage.tsx),
    # distinguishable from a sub-criterion the user never touched at all
    # (which never appears in the payload in the first place).
    assert svc.normalize_scores({"written_1": {"score": None, "notes": "  "}}) == {"written_1": {"score": None, "notes": None}}
    assert svc.normalize_scores({"written_1": {"score": None, "notes": " ok "}}) == {"written_1": {"score": None, "notes": "ok"}}


def test_manager_evaluation_scores_averages_across_reviewers():
    pooled = svc.manager_evaluation_scores([
        _scores({"written_1": 4, "written_2": 2}),
        _scores({"written_1": 5}),  # written_2 untouched by this reviewer — excluded, not counted as 0
    ])
    assert pooled["written_1"]["score"] == 4.5
    assert pooled["written_2"]["score"] == 2.0


def test_create_defaults_reviewer_hours_description_and_format(db):
    review = svc.create_review(db, PerformanceReviewCreate(project_id="prj", employee_id="emp", review_date=date(2025, 12, 15)), "mgr")
    assert _reviewer_ids(review) == ["mgr"]  # falls back to the project manager
    assert float(review.duration_hours) == 296
    assert review.project_description == "Cash flow model"
    assert review.status == "self_assessment"
    assert review.evaluation_format == "standard"
    assert review.self_due_date is None

    scoped = svc.create_review(db, PerformanceReviewCreate(
        project_id="prj", employee_id="emp", review_date=date(2025, 12, 15),
        period_start=date(2025, 6, 1), period_end=date(2025, 12, 31),
        self_due_date=date(2026, 1, 1),
    ), "mgr")
    assert float(scoped.duration_hours) == 96
    assert scoped.self_due_date == date(2026, 1, 1)


def test_create_with_multiple_reviewers(db):
    review = svc.create_review(db, PerformanceReviewCreate(
        project_id="prj", employee_id="emp", review_date=date(2025, 12, 15),
        reviewer_ids=["mgr", "mgr2"],
    ), "mgr")
    assert set(_reviewer_ids(review)) == {"mgr", "mgr2"}
    # The employee themselves is silently dropped if somehow included.
    review2 = svc.create_review(db, PerformanceReviewCreate(
        project_id="prj", employee_id="emp", review_date=date(2025, 12, 15),
        reviewer_ids=["mgr", "emp"],
    ), "mgr")
    assert _reviewer_ids(review2) == ["mgr"]


def test_evaluation_formats_registry():
    assert svc.EVALUATION_FORMATS["standard"] == "Impact Point Standard"


def _visible(db, review, viewer, manage=False):
    out = svc.serialize(db, [review], viewer, manage)[0]
    vis = {
        "self": out["evaluations"]["self"]["visible"],
        "joint": out["evaluations"]["joint"]["visible"],
        "reviewers": {r["reviewer_id"]: r["visible"] for r in out["evaluations"]["reviewers"]},
    }
    return vis, out


def test_self_is_blind_to_everyone_but_the_reviewee_until_final_submission(db):
    review = svc.create_review(db, PerformanceReviewCreate(project_id="prj", employee_id="emp", review_date=date(2025, 12, 15), reviewer_ids=["mgr"]), "x")
    svc.update_review(db, review, {"self_section_comments": {"written": "Did well here."}, "self_strengths": "Excel"})
    vis, out = _visible(db, review, "emp")
    assert vis == {"self": True, "joint": False, "reviewers": {"mgr": False}}
    assert out["evaluations"]["self"]["section_comments"] == {"written": "Did well here."}
    vis, out = _visible(db, review, "mgr")
    assert vis == {"self": False, "joint": False, "reviewers": {"mgr": True}}
    assert out["self_strengths"] is None  # not visible yet

    svc.transition_review(db, review, "submit_self")
    svc.apply_reviewer_evaluation(db, review, "mgr", {"scores": _scores({"written_1": 3, "verbal_1": 4})})
    db.commit()
    # Self is submitted but reviewers STILL can't see its content — only at
    # Final Submission (joint_review), per the explicit "no antes" requirement.
    vis, out = _visible(db, review, "mgr")
    assert vis["self"] is False and out["self_strengths"] is None
    vis, out = _visible(db, review, "emp")
    assert vis["reviewers"]["mgr"] is False
    assert out["overall_average"] is None

    review = svc.transition_review(db, review, "submit_manager", actor_id="mgr")
    assert review.status == "joint_review"
    vis, out = _visible(db, review, "emp")
    assert vis["self"] is True and out["self_strengths"] == "Excel"
    assert out["overall_average"] == out["evaluations"]["joint"]["overall_average"]


def test_reviewer_cannot_start_before_self_is_submitted(db):
    review = svc.create_review(db, PerformanceReviewCreate(project_id="prj", employee_id="emp", review_date=date(2025, 12, 15), reviewer_ids=["mgr"]), "x")
    # Self hasn't been submitted yet — the reviewer has no editable fields at all.
    assert svc.editable_fields(review, "mgr", False) == set()
    with pytest.raises(ValueError):
        svc.transition_review(db, review, "submit_manager", actor_id="mgr")

    svc.transition_review(db, review, "submit_self")
    assert svc.editable_fields(review, "mgr", False) == {"reviewer_evaluation"}


def test_employee_never_sees_any_individual_reviewer_row(db):
    review = svc.create_review(db, PerformanceReviewCreate(
        project_id="prj", employee_id="emp", review_date=date(2025, 12, 15), reviewer_ids=["mgr", "mgr2"],
    ), "x")
    svc.transition_review(db, review, "submit_self")
    svc.apply_reviewer_evaluation(db, review, "mgr", {"scores": _scores({"written_1": 4})})
    svc.apply_reviewer_evaluation(db, review, "mgr2", {"scores": _scores({"written_1": 2})})
    db.commit()
    svc.transition_review(db, review, "submit_manager", actor_id="mgr")
    review = svc.transition_review(db, review, "submit_manager", actor_id="mgr2")
    assert review.status == "joint_review"

    # Reviewers see each other now that they're "together"...
    vis, _ = _visible(db, review, "mgr")
    assert vis["reviewers"] == {"mgr": True, "mgr2": True}
    # ...but the reviewee never sees either individual row, only the joint result.
    vis, out = _visible(db, review, "emp")
    assert vis["reviewers"] == {"mgr": False, "mgr2": False}
    assert all(r["scores"] == {} for r in out["evaluations"]["reviewers"])
    assert vis["joint"] is True  # the official result stays visible to them


def test_self_assessment_deadline_locks_editing(db):
    yesterday = date.today() - timedelta(days=1)
    tomorrow = date.today() + timedelta(days=1)
    locked = svc.create_review(db, PerformanceReviewCreate(
        project_id="prj", employee_id="emp", review_date=date(2025, 12, 15), self_due_date=yesterday,
    ), "x")
    assert svc.self_assessment_locked(locked) is True
    assert svc.editable_fields(locked, "emp", False) == set()
    # Admin/Manager always bypasses the deadline.
    assert "self_section_comments" in svc.editable_fields(locked, "anyone", True)

    open_review = svc.create_review(db, PerformanceReviewCreate(
        project_id="prj", employee_id="emp", review_date=date(2025, 12, 15), self_due_date=tomorrow,
    ), "x")
    assert svc.self_assessment_locked(open_review) is False
    assert "self_section_comments" in svc.editable_fields(open_review, "emp", False)

    no_deadline = svc.create_review(db, PerformanceReviewCreate(project_id="prj", employee_id="emp", review_date=date(2025, 12, 15)), "x")
    assert svc.self_assessment_locked(no_deadline) is False


def test_joint_only_opens_once_every_reviewer_submits(db):
    review = svc.create_review(db, PerformanceReviewCreate(
        project_id="prj", employee_id="emp", review_date=date(2025, 12, 15), reviewer_ids=["mgr", "mgr2"],
    ), "x")
    svc.transition_review(db, review, "submit_self")
    svc.apply_reviewer_evaluation(db, review, "mgr", {"scores": _scores({"written_1": 4})})
    svc.apply_reviewer_evaluation(db, review, "mgr2", {"scores": _scores({"written_1": 2})})
    db.commit()

    # mgr2 is blind to mgr's row and vice versa, even though both are reviewers.
    vis, _ = _visible(db, review, "mgr")
    assert vis["reviewers"] == {"mgr": True, "mgr2": False}

    review = svc.transition_review(db, review, "submit_manager", actor_id="mgr")
    assert review.status == "in_review"  # mgr2 hasn't submitted yet
    assert review.joint_scores == {}

    review = svc.transition_review(db, review, "submit_manager", actor_id="mgr2")
    assert review.status == "joint_review"
    # Joint pre-fills from the average of both reviewers, rounded to a whole score.
    assert review.joint_scores["written_1"]["score"] == 3  # mean(4, 2) = 3.0

    with pytest.raises(ValueError):
        svc.transition_review(db, review, "submit_manager", actor_id="mgr")  # already submitted


def test_reopen_from_joint_preserves_reviewer_submissions(db):
    review = svc.create_review(db, PerformanceReviewCreate(
        project_id="prj", employee_id="emp", review_date=date(2025, 12, 15), reviewer_ids=["mgr"],
    ), "x")
    review = _submit_self_and_reviewer(db, review, "mgr", {"written_1": 4})
    assert review.status == "joint_review"

    review = svc.transition_review(db, review, "reopen")
    assert review.status == "in_review"
    assert review.reviewers[0].submitted_at is not None  # preserved, not cleared
    assert review.reviewers[0].scores["written_1"]["score"] == 4


def test_editable_fields_follow_the_stage(db):
    review = svc.create_review(db, PerformanceReviewCreate(project_id="prj", employee_id="emp", review_date=date(2025, 12, 15), reviewer_ids=["mgr"]), "x")
    assert "self_section_comments" in svc.editable_fields(review, "emp", False)
    assert svc.editable_fields(review, "mgr", False) == set()  # can't start until self is submitted
    svc.transition_review(db, review, "submit_self")
    assert svc.editable_fields(review, "emp", False) == set()
    assert svc.editable_fields(review, "mgr", False) == {"reviewer_evaluation"}
    svc.apply_reviewer_evaluation(db, review, "mgr", {"scores": _scores({"written_1": 3})})
    db.commit()
    svc.transition_review(db, review, "submit_manager", actor_id="mgr")
    assert svc.editable_fields(review, "mgr", False) == svc.JOINT_FIELDS
    svc.transition_review(db, review, "complete")
    assert svc.editable_fields(review, "mgr", False) == set()
    assert svc.editable_fields(review, "anyone", True) == svc.ADMIN_FIELDS


def test_analytics_uses_completed_reviews_only(db):
    done = svc.create_review(db, PerformanceReviewCreate(project_id="prj", employee_id="emp", review_date=date(2025, 12, 15), reviewer_ids=["mgr"]), "x")
    svc.update_review(db, done, {"joint_scores": _scores({"written_1": 4})})
    svc.apply_reviewer_evaluation(db, done, "mgr", {"scores": _scores({"written_1": 3})})
    done.status = "completed"
    db.commit()
    svc.create_review(db, PerformanceReviewCreate(project_id="prj", employee_id="emp", review_date=date(2026, 1, 15)), "x")

    data = svc.analytics(db, 2025)
    assert data["years"] == [2025]
    assert len(data["reviews"]) == 1
    row = data["reviews"][0]
    # Self no longer scores anything — "self" overall is always None now.
    assert (row["self"]["overall"], row["manager"]["overall"], row["joint"]["overall"]) == (None, 3.0, 4.0)
    assert row["joint"]["items"] == {"written_1": 4}
    assert row["reviewer_names"] == ["Brandon Beal"]
    assert svc.analytics(db, 2026)["reviews"] == []


def test_xlsx_export_layout_single_reviewer(db):
    review = svc.create_review(db, PerformanceReviewCreate(project_id="prj", employee_id="emp", review_date=date(2025, 12, 15), reviewer_ids=["mgr"]), "mgr")
    svc.update_review(db, review, {
        "joint_scores": _scores(SAMPLE_SCORES),
        "self_strengths": "Excel skills",
        "self_section_comments": {"written": "Clear writer."},
        "joint_notes": "Agreed",
    })
    svc.apply_reviewer_evaluation(db, review, "mgr", {"scores": _scores({"written_1": 2}), "strengths_notes": "Technically sound"})
    db.commit()
    data = svc.review_export_data(db, review)
    wb = load_workbook(BytesIO(generate_performance_review_xlsx(data)))
    assert wb.sheetnames == ["Final Submission", "Manager", "Self"]  # single reviewer: no extra per-reviewer sheet
    ws = wb["Final Submission"]

    assert export_filename(data) == "THERAPAK, LLC - Michael Franz - 2025.xlsx"
    assert ws["B6"].value == "Franz, Michael"
    assert ws["B8"].value == "Brandon Beal"
    assert ws["B11"].value == "296 hours"
    assert ws["B12"].value == "12/15/2025"
    assert [ws[f"D{r}"].value for r in range(7, 13)] == [3.0, 3.25, 2.8, 3.0, 2.75, 2.96]
    assert ws["B14"].value == "Excel skills" and ws["D14"].value is None  # Final Submission sheet shows no single reviewer's notes
    assert ws["B17"].value == "Agreed"
    assert ws["A19"].value == "Written Communications" and ws["C19"].value == 3.0
    assert ws["C20"].value is None and ws["C21"].value == 3
    assert ws["A43"].value == "Overall Professionalism and Core Values" and ws["C43"].value == 2.75
    assert ws["C44"].value == 2

    assert wb["Manager"]["C20"].value == 2 and wb["Manager"]["D12"].value == 2.0

    self_ws = wb["Self"]
    # No per-item scores at all on the Self sheet — just the reference list…
    assert self_ws["A19"].value == "Written Communications" and self_ws["C19"].value is None
    assert self_ws["B20"].value and self_ws["C20"].value is None  # sub-criterion label, no score column
    # …and one comment per section, written in the criterion's own row.
    assert self_ws["D19"].value == "Clear writer."


def test_xlsx_export_adds_a_sheet_per_reviewer_when_there_are_several(db):
    review = svc.create_review(db, PerformanceReviewCreate(
        project_id="prj", employee_id="emp", review_date=date(2025, 12, 15), reviewer_ids=["mgr", "mgr2"],
    ), "mgr")
    svc.apply_reviewer_evaluation(db, review, "mgr", {"scores": _scores({"written_1": 4}), "strengths_notes": "From mgr"})
    svc.apply_reviewer_evaluation(db, review, "mgr2", {"scores": _scores({"written_1": 2}), "strengths_notes": "From mgr2"})
    db.commit()
    data = svc.review_export_data(db, review)
    wb = load_workbook(BytesIO(generate_performance_review_xlsx(data)))
    assert wb.sheetnames == ["Final Submission", "Manager", "Self", "Brandon Beal", "Casey Diaz"]
    assert wb["Manager"]["C20"].value == 3.0  # mean(4, 2)
    assert wb["Brandon Beal"]["C20"].value == 4 and wb["Brandon Beal"]["D14"].value == "From mgr"
    assert wb["Casey Diaz"]["C20"].value == 2 and wb["Casey Diaz"]["D14"].value == "From mgr2"


def test_reviewer_never_defaults_to_the_reviewee(db):
    # The project manager being reviewed falls through to the owner (none here).
    review = svc.create_review(db, PerformanceReviewCreate(project_id="prj", employee_id="mgr", review_date=date(2025, 12, 15)), "mgr")
    assert review.reviewers == []


def test_projects_panel_and_team(db):
    from models.employee_projects import EmployeeProject
    db.add(EmployeeProject(id="ep1", user_id="emp", project_id="prj"))
    db.add(Project(id="int", client_id="cli", name="Internal", is_internal=True))
    db.commit()

    panel = svc.list_review_projects(db)
    assert [p["id"] for p in panel] == ["prj"]  # internal projects excluded
    assert panel[0]["performance_review_enabled"] is False and panel[0]["team_size"] == 1

    svc.set_project_review_enabled(db, db.get(Project, "prj"), True)
    svc.create_review(db, PerformanceReviewCreate(project_id="prj", employee_id="emp", review_date=date(2025, 12, 15), reviewer_ids=["mgr"]), "mgr")
    panel = svc.list_review_projects(db)
    assert panel[0]["performance_review_enabled"] is True
    assert (panel[0]["reviews_total"], panel[0]["reviews_self_assessment"]) == (1, 1)

    team = {m["employee_id"]: m for m in svc.project_team(db, "prj")}
    # Assigned employee + someone who only logged time — mgr is also the
    # reviewer, but that alone doesn't add them to the roster (they're
    # already here via logged hours).
    assert set(team) == {"emp", "mgr"}
    assert team["emp"]["is_assigned"] and not team["mgr"]["is_assigned"]
    assert team["emp"]["logged_hours"] == 296 and team["mgr"]["logged_hours"] == 50
    assert len(team["emp"]["reviews"]) == 1 and team["emp"]["reviews"][0]["reviewer_names"] == ["Brandon Beal"]
    assert team["mgr"]["reviews"] == []
    assert svc.has_open_review(db, "prj", "emp") and not svc.has_open_review(db, "prj", "mgr")
